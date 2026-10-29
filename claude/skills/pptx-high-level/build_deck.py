#!/usr/bin/env python3
"""Render a concise, high-level PowerPoint from a JSON deck spec.

Usage:
    python3 build_deck.py <deck.json> <out.pptx>

The JSON spec is authored by the skill. Keep it high-level and short, and
prefer diagram slides (flow / compare / stack / two_column / mermaid) over walls
of text. See SKILL.md for the slide grammar. This script only handles the
layout/theme; it never invents content.

Theme: "modern professional" -- crisp sans, a deep-teal accent, soft rounded
cards with a subtle shadow and light tint, generous whitespace, and a quiet
footer (deck title + slide number) on content slides. Tuned for research
meetings and conference-style talks.
"""
import json
import os
import shutil
import subprocess
import sys
import tempfile

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

# 16:9 canvas
EMU_W, EMU_H = Inches(13.333), Inches(7.5)

# ---- theme: modern professional, deep-teal accent ---------------------------
INK = RGBColor(0x17, 0x1B, 0x22)      # near-black body text (cool)
MUTED = RGBColor(0x5B, 0x64, 0x70)    # secondary / footer text
FAINT = RGBColor(0xE2, 0xE6, 0xEA)    # hairlines
ACCENT = RGBColor(0x0F, 0x76, 0x6E)   # deep teal: accents, titles, highlights
ACCENT_DK = RGBColor(0x0B, 0x59, 0x53)  # darker teal for text on light
PAPER = RGBColor(0xFF, 0xFF, 0xFF)    # slide background
PANEL = RGBColor(0xF4, 0xF6, 0xF8)    # neutral card fill
NODE = RGBColor(0xE4, 0xF1, 0xEF)     # teal-tinted card fill
SUMMARY = RGBColor(0xEA, 0xF4, 0xF2)  # summary panel fill
BAR = RGBColor(0xD4, 0xDA, 0xDF)      # non-highlighted bar
GOOD = RGBColor(0x1B, 0x7F, 0x3B)     # "after" green / kept / done
WARN = RGBColor(0xB0, 0x45, 0x1E)     # "before" rust / dropped
AMBER = RGBColor(0xB4, 0x7A, 0x0F)    # "testing" / in-progress amber
DARK = RGBColor(0x15, 0x1B, 0x24)     # closing-slide background

# tints for status pills / tiles (light fills behind the status color)
GOOD_TINT = RGBColor(0xEC, 0xF6, 0xEF)
WARN_TINT = RGBColor(0xFB, 0xEE, 0xE8)
AMBER_TINT = RGBColor(0xFA, 0xF3, 0xE2)


def _status_colors(status):
    """Map a status word to (solid, tint, label). Used by status_list / timeline.
    kept/done/good/win -> green; dropped/cut/bad/no -> rust; else -> amber."""
    s = (status or "").strip().lower()
    if s in ("kept", "done", "good", "win", "yes", "shipped", "keep"):
        return GOOD, GOOD_TINT, (status or "kept").upper()
    if s in ("dropped", "cut", "bad", "no", "failed", "drop", "regressed"):
        return WARN, WARN_TINT, (status or "dropped").upper()
    return AMBER, AMBER_TINT, (status or "testing").upper()

FONT = "Calibri"

# mermaid node/line colors (hex strings; recolored by the palette)
MM_NODE = "#E4F1EF"
MM_BORDER = "#0F766E"
MM_TEXT = "#171B22"
MM_LINE = "#0F766E"

# Optional palette override, chosen via top-level spec "theme". Keeps the default
# teal but lets a deck pick a different accent so it does not feel same-y.
PALETTES = {
    "teal": {},
    "indigo": {"ACCENT": (0x43, 0x38, 0xCA), "ACCENT_DK": (0x31, 0x2E, 0x81),
               "NODE": (0xEC, 0xEC, 0xFB), "SUMMARY": (0xEE, 0xEE, 0xFC),
               "MM_NODE": "#ECECFB", "MM_BORDER": "#4338CA", "MM_LINE": "#4338CA"},
    "navy": {"ACCENT": (0x1E, 0x3A, 0x8A), "ACCENT_DK": (0x17, 0x2C, 0x66),
             "NODE": (0xE7, 0xED, 0xF9), "SUMMARY": (0xEC, 0xF1, 0xFB),
             "MM_NODE": "#E7EDF9", "MM_BORDER": "#1E3A8A", "MM_LINE": "#1E3A8A"},
    "plum": {"ACCENT": (0x7C, 0x3A, 0x6A), "ACCENT_DK": (0x5C, 0x2B, 0x4F),
             "NODE": (0xF3, 0xE9, 0xF0), "SUMMARY": (0xF6, 0xED, 0xF3),
             "MM_NODE": "#F3E9F0", "MM_BORDER": "#7C3A6A", "MM_LINE": "#7C3A6A"},
    "slate": {"ACCENT": (0x33, 0x41, 0x55), "ACCENT_DK": (0x24, 0x2F, 0x40),
              "NODE": (0xEC, 0xEF, 0xF3), "SUMMARY": (0xEF, 0xF2, 0xF6),
              "MM_NODE": "#ECEFF3", "MM_BORDER": "#334155", "MM_LINE": "#334155"},
}


def apply_palette(name):
    """Recolor the theme accent from a named palette (default 'teal' = no change)."""
    p = PALETTES.get((name or "teal").lower())
    if not p:
        return
    g = globals()
    for k, v in p.items():
        g[k] = v if isinstance(v, str) else RGBColor(*v)


# content region below the title; footer lives under this
CONTENT_BOTTOM = Inches(6.55)


def _blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _bg(slide, color=PAPER):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def _soft_shadow(shape, alpha=26000):
    """Inject a subtle outer shadow (the 'gentle depth' of the card look)."""
    spPr = shape._element.spPr
    old = spPr.find(qn("a:effectLst"))
    if old is not None:
        spPr.remove(old)
    eff = spPr.makeelement(qn("a:effectLst"), {})
    shdw = eff.makeelement(qn("a:outerShdw"), {
        "blurRad": str(int(Pt(9))),
        "dist": str(int(Pt(2.5))),
        "dir": "5400000",   # straight down
        "rotWithShape": "0",
    })
    clr = shdw.makeelement(qn("a:srgbClr"), {"val": "10202C"})
    a = clr.makeelement(qn("a:alpha"), {"val": str(alpha)})
    clr.append(a)
    shdw.append(clr)
    eff.append(shdw)
    spPr.append(eff)


def _tb(slide, left, top, width, height, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    return tf


def _para(tf, text, size, color=INK, bold=False, align=PP_ALIGN.LEFT,
          bullet=False, space_after=10, first=False):
    p = tf.paragraphs[0] if first and not tf.paragraphs[0].runs else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    if bullet:
        p.level = 0
    r = p.add_run()
    r.text = ("-  " + text) if bullet else text
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return p


def _rect(slide, shape, left, top, width, height, fill, line=None, shadow=False):
    box = slide.shapes.add_shape(shape, int(left), int(top), int(width), int(height))
    box.fill.solid()
    box.fill.fore_color.rgb = fill
    if line is None:
        box.line.fill.background()
    else:
        box.line.color.rgb = line
        box.line.width = Pt(1)
    box.shadow.inherit = False
    if shadow:
        _soft_shadow(box)
    return box


def _panel(slide, left, top, width, height, fill=None, shadow=True):
    return _rect(slide, MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height,
                 PANEL if fill is None else fill, line=None, shadow=shadow)


def _node(slide, left, top, width, height, text, fill=None, text_color=INK,
          size=18, bold=True, shape=MSO_SHAPE.ROUNDED_RECTANGLE, shadow=True):
    """A soft labeled card: the atom of every diagram slide."""
    if fill is None:
        fill = NODE
    box = _rect(slide, shape, left, top, width, height, fill, line=None, shadow=shadow)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.12)
    tf.margin_right = Inches(0.12)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = text_color
    return box


def _title_line(slide, text):
    """A refined slide title: a short teal accent tick above dark-ink text."""
    _rect(slide, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.72), Inches(0.6),
          Inches(0.52), Inches(0.11), ACCENT)
    tf = _tb(slide, Inches(0.7), Inches(0.82), Inches(12), Inches(0.95))
    _para(tf, text, 29, INK, bold=True, first=True)


def _caption(slide, text):
    if not text:
        return
    tf = _tb(slide, Inches(0.7), Inches(6.62), Inches(12), Inches(0.4))
    _para(tf, text, 15, MUTED, align=PP_ALIGN.CENTER, first=True, space_after=0)


def _footer(slide, ctx):
    """Quiet footer: deck title on the left, 'n / total' on the right."""
    if not ctx:
        return
    # hairline separating the footer from the content
    _rect(slide, MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(7.0),
          Inches(11.93), Pt(0.75), FAINT)
    if ctx.get("title"):
        tf = _tb(slide, Inches(0.7), Inches(7.08), Inches(9.5), Inches(0.35))
        _para(tf, ctx["title"], 10.5, MUTED, first=True, space_after=0)
    tf = _tb(slide, Inches(10.7), Inches(7.08), Inches(1.93), Inches(0.35))
    _para(tf, "%d / %d" % (ctx["index"], ctx["total"]), 10.5, MUTED,
          align=PP_ALIGN.RIGHT, first=True, space_after=0)


# ---- mermaid rendering -----------------------------------------------------

_MMDC = shutil.which("mmdc") or os.path.expanduser("~/.local/bin/mmdc")
_CHROME = (shutil.which("google-chrome") or shutil.which("chromium")
           or shutil.which("brave-browser") or shutil.which("chromium-browser"))


def _render_mermaid(src, scale=3):
    """Render mermaid source to a PNG. Returns the path, or None on failure."""
    if not _MMDC or not os.path.exists(_MMDC):
        return None
    d = tempfile.mkdtemp(prefix="pptx_mmd_")
    mmd, png = os.path.join(d, "d.mmd"), os.path.join(d, "d.png")
    pupp, cfg = os.path.join(d, "pupp.json"), os.path.join(d, "cfg.json")
    with open(mmd, "w") as f:
        f.write(src)
    with open(pupp, "w") as f:
        json.dump({"executablePath": _CHROME, "args": ["--no-sandbox"]}
                  if _CHROME else {"args": ["--no-sandbox"]}, f)
    with open(cfg, "w") as f:
        json.dump({"theme": "base",
                   "themeVariables": {"fontFamily": FONT + ", Arial, sans-serif",
                                      "primaryColor": MM_NODE,
                                      "primaryBorderColor": MM_BORDER,
                                      "primaryTextColor": MM_TEXT,
                                      "secondaryColor": "#F4F6F8",
                                      "tertiaryColor": "#F4F6F8",
                                      "lineColor": MM_LINE},
                   "flowchart": {"htmlLabels": True, "curve": "basis"}}, f)
    env = dict(os.environ)
    if _CHROME:
        env["PUPPETEER_EXECUTABLE_PATH"] = _CHROME
    cmd = [_MMDC, "-i", mmd, "-o", png, "-b", "transparent", "-s", str(scale),
           "-p", pupp, "-c", cfg]
    try:
        subprocess.run(cmd, env=env, capture_output=True, timeout=150, check=True)
    except Exception:
        return None
    return png if os.path.exists(png) else None


def _png_size(path):
    """Return (w, h) of a PNG from its header, no dependencies."""
    try:
        with open(path, "rb") as f:
            head = f.read(24)
        if head[:8] == b"\x89PNG\r\n\x1a\n":
            return int.from_bytes(head[16:20], "big"), int.from_bytes(head[20:24], "big")
    except Exception:
        pass
    return (1600, 900)


# ---- text slide renderers --------------------------------------------------

def slide_title(prs, s):
    slide = _blank(prs)
    _bg(slide)
    # vertical teal accent rule to the left of the title block
    _rect(slide, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(2.55),
          Inches(0.14), Inches(2.25), ACCENT)
    tf = _tb(slide, Inches(1.28), Inches(2.4), Inches(11.2), Inches(2.6),
             anchor=MSO_ANCHOR.MIDDLE)
    _para(tf, s["title"], 44, INK, bold=True, first=True, space_after=8)
    if s.get("subtitle"):
        _para(tf, s["subtitle"], 22, MUTED, space_after=0)
    return slide


def slide_summary(prs, s):
    slide = _blank(prs)
    _bg(slide)
    _title_line(slide, s.get("title", "In one sentence"))
    _panel(slide, Inches(0.7), Inches(2.3), Inches(11.93), Inches(2.9),
           fill=SUMMARY)
    # inner accent bar on the left edge of the panel
    _rect(slide, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.95), Inches(2.6),
          Inches(0.1), Inches(2.3), ACCENT)
    tf = _tb(slide, Inches(1.35), Inches(2.3), Inches(10.9), Inches(2.9),
             anchor=MSO_ANCHOR.MIDDLE)
    _para(tf, s["body"], 27, INK, bold=True, first=True)
    return slide


def slide_section(prs, s):
    slide = _blank(prs)
    _bg(slide, ACCENT)
    tf = _tb(slide, Inches(0.9), 0, Inches(11.5), EMU_H, anchor=MSO_ANCHOR.MIDDLE)
    _para(tf, s["title"], 40, PAPER, bold=True, first=True, space_after=8)
    if s.get("subtitle"):
        _para(tf, s["subtitle"], 22, RGBColor(0xC9, 0xE8, 0xE4), space_after=0)
    return slide


def slide_bullets(prs, s):
    slide = _blank(prs)
    _bg(slide)
    _title_line(slide, s["title"])
    tf = _tb(slide, Inches(0.9), Inches(2.0), Inches(11.5), Inches(4.4))
    bullets = s.get("bullets", [])
    for i, b in enumerate(bullets):
        _para(tf, b, 22, INK, bullet=True, space_after=16, first=(i == 0))
    return slide


def slide_two_column(prs, s):
    slide = _blank(prs)
    _bg(slide)
    _title_line(slide, s.get("title", "Before vs After"))
    cols = [
        ("left", Inches(0.7), s.get("left_title", "Before"), WARN),
        ("right", Inches(6.93), s.get("right_title", "After"), GOOD),
    ]
    for key, x, head, color in cols:
        _panel(slide, x, Inches(2.0), Inches(5.7), Inches(4.5))
        # colored accent tick per column header
        _rect(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x + Inches(0.3), Inches(2.28),
              Inches(0.4), Inches(0.09), color)
        htf = _tb(slide, x + Inches(0.3), Inches(2.44), Inches(5.1), Inches(0.6))
        _para(htf, head, 21, color, bold=True, first=True)
        btf = _tb(slide, x + Inches(0.3), Inches(3.15), Inches(5.1), Inches(3.2))
        items = s.get(key, [])
        for i, b in enumerate(items):
            _para(btf, b, 18, INK, bullet=True, space_after=12, first=(i == 0))
    return slide


def slide_closing(prs, s):
    slide = _blank(prs)
    _bg(slide, DARK)
    _rect(slide, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(3.05),
          Inches(0.52), Inches(0.11), ACCENT)
    tf = _tb(slide, Inches(0.9), Inches(3.3), Inches(11.5), Inches(2.2))
    _para(tf, s.get("title", "Takeaway"), 36, PAPER, bold=True, first=True,
          space_after=8)
    if s.get("body"):
        _para(tf, s["body"], 22, RGBColor(0xC9, 0xD1, 0xDB), space_after=0)
    return slide


# ---- diagram slide renderers -----------------------------------------------

def slide_flow(prs, s):
    """A left-to-right (or top-down) sequence of soft cards joined by arrows.
    Spec: steps: [str, ...] (<=6). Optional direction: "right" | "down", caption."""
    slide = _blank(prs)
    _bg(slide)
    _title_line(slide, s["title"])
    steps = [str(x) for x in s.get("steps", [])][:6]
    n = len(steps)
    if n == 0:
        return slide
    down = s.get("direction") == "down"
    if down:
        margin_x, top = Inches(3.5), Inches(2.0)
        avail_h = int(CONTENT_BOTTOM - top)
        arrow_h = Inches(0.42)
        box_w = int(EMU_W - margin_x * 2)
        box_h = int((avail_h - arrow_h * (n - 1)) / n)
        y = int(top)
        for i, step in enumerate(steps):
            _node(slide, int(margin_x), y, box_w, box_h, step, size=20)
            y += box_h
            if i < n - 1:
                _rect(slide, MSO_SHAPE.DOWN_ARROW,
                      int(EMU_W / 2 - Inches(0.32)), y,
                      int(Inches(0.64)), int(arrow_h), ACCENT)
                y += int(arrow_h)
    else:
        margin, top = Inches(0.6), Inches(3.0)
        box_h = Inches(1.7)
        arrow_w = Inches(0.55)
        total_w = int(EMU_W - margin * 2)
        box_w = int((total_w - int(arrow_w) * (n - 1)) / n)
        x = int(margin)
        for i, step in enumerate(steps):
            _node(slide, x, int(top), box_w, int(box_h), step, size=19)
            x += box_w
            if i < n - 1:
                _rect(slide, MSO_SHAPE.RIGHT_ARROW,
                      x, int(top + box_h / 2 - Inches(0.26)),
                      int(arrow_w), int(Inches(0.52)), ACCENT)
                x += int(arrow_w)
    _caption(slide, s.get("caption"))
    return slide


_BAR_ROLES = {"good": (GOOD, GOOD), "warn": (WARN, WARN), "accent": (ACCENT, ACCENT_DK),
              "muted": (BAR, INK)}


def slide_compare(prs, s):
    """Horizontal bars comparing a few labeled values (benchmark numbers, etc.).
    Spec: bars: [{label, value}, ...] (<=6). Optional unit (e.g. "%"), highlight.
    Per-bar color: give a bar a "role" of "good" | "warn" | "accent" | "muted"
    (e.g. mark ours "accent", the best competitor "good", the rest "muted") for a
    clear multi-way comparison. `highlight` (one label) still works as a shortcut."""
    slide = _blank(prs)
    _bg(slide)
    _title_line(slide, s["title"])
    bars = s.get("bars", [])[:6]
    if not bars:
        return slide
    unit = s.get("unit", "")
    highlight = s.get("highlight")
    maxv = max([float(b.get("value", 0)) for b in bars] + [1.0])
    top = Inches(2.15)
    label_w = Inches(3.2)
    bar_x = int(Inches(0.7) + label_w + Inches(0.2))
    bar_max_w = int(EMU_W - bar_x - Inches(1.7))
    n = len(bars)
    row_h = int((int(CONTENT_BOTTOM) - int(top)) / n)
    bar_h = min(int(Inches(0.62)), int(row_h * 0.58))
    for i, b in enumerate(bars):
        y = int(top) + row_h * i
        val = float(b.get("value", 0))
        role = b.get("role")
        is_hi = highlight is not None and b.get("label") == highlight
        if role in _BAR_ROLES:
            fill, vcolor = _BAR_ROLES[role]
            emphasized = role != "muted"
        elif is_hi:
            fill, vcolor, emphasized = ACCENT, ACCENT_DK, True
        else:
            fill, vcolor, emphasized = BAR, INK, False
        ltf = _tb(slide, Inches(0.7), y, label_w, row_h, anchor=MSO_ANCHOR.MIDDLE)
        _para(ltf, str(b.get("label", "")), 18, INK, bold=emphasized,
              align=PP_ALIGN.RIGHT, first=True)
        w = max(int(Inches(0.05)), int(bar_max_w * val / maxv))
        by = y + (row_h - bar_h) // 2
        _rect(slide, MSO_SHAPE.ROUNDED_RECTANGLE, bar_x, by, w, bar_h,
              fill, shadow=emphasized)
        vtf = _tb(slide, bar_x + w + Inches(0.12), y, Inches(1.5), row_h,
                  anchor=MSO_ANCHOR.MIDDLE)
        _para(vtf, "%s%s" % (b.get("value"), unit), 18,
              vcolor, bold=True, first=True)
    _caption(slide, s.get("caption"))
    return slide


def slide_stack(prs, s):
    """Vertically stacked full-width cards (layers / architecture, top to bottom).
    Spec: layers: [str, ...] (<=5). Optional caption."""
    slide = _blank(prs)
    _bg(slide)
    _title_line(slide, s["title"])
    layers = [str(x) for x in s.get("layers", [])][:5]
    n = len(layers)
    if n == 0:
        return slide
    top = Inches(2.1)
    margin = Inches(1.6)
    gap = Inches(0.24)
    avail_h = int(CONTENT_BOTTOM - int(top))
    box_h = int((avail_h - int(gap) * (n - 1)) / n)
    w = int(EMU_W - int(margin) * 2)
    for i, layer in enumerate(layers):
        y = int(top) + (box_h + int(gap)) * i
        fill = NODE if i % 2 == 0 else PANEL
        _node(slide, int(margin), y, w, box_h, layer, fill=fill, size=20)
    _caption(slide, s.get("caption"))
    return slide


def slide_mermaid(prs, s):
    """A real Mermaid diagram, rasterized and embedded (best-looking option).
    Spec: mermaid: "<mermaid source>". Optional caption.
    Optional explain: [str, ...] -- plain-language points shown in a side panel so
    the diagram is self-explanatory (nothing ambiguous). With explain, the diagram
    sits on the left and the explanation card on the right.
    Keep it high-level: a handful of nodes, short labels."""
    slide = _blank(prs)
    _bg(slide)
    _title_line(slide, s["title"])
    explain = [e for e in (s.get("explain") or []) if e]
    cy, ch = int(Inches(1.9)), int(Inches(4.7))
    if explain:
        cx, cw = int(Inches(0.7)), int(Inches(8.0))
        ex_left, ex_w = int(Inches(9.0)), int(Inches(3.63))
    else:
        cx, cw = int(Inches(0.7)), int(Inches(12.0))
    png = _render_mermaid(s.get("mermaid", ""))
    if png:
        iw, ih = _png_size(png)
        ar = (iw / ih) if ih else 16 / 9
        w = cw
        h = int(w / ar)
        if h > ch:
            h = ch
            w = int(h * ar)
        left = cx + (cw - w) // 2
        top = cy + (ch - h) // 2
        slide.shapes.add_picture(png, left, top, width=w, height=h)
    else:
        tf = _tb(slide, cx, cy, cw, ch, anchor=MSO_ANCHOR.MIDDLE)
        _para(tf, "(mermaid render unavailable: install @mermaid-js/mermaid-cli)",
              20, MUTED, align=PP_ALIGN.CENTER, first=True)
    if explain:
        _panel(slide, ex_left, cy, ex_w, ch, fill=PANEL, shadow=True)
        _rect(slide, MSO_SHAPE.ROUNDED_RECTANGLE, ex_left + Inches(0.28), cy + Inches(0.28),
              Inches(0.4), Inches(0.09), ACCENT)
        tf = _tb(slide, ex_left + Inches(0.28), cy + Inches(0.44),
                 ex_w - Inches(0.56), ch - Inches(0.6))
        _para(tf, "In plain terms", 14, ACCENT_DK, bold=True, first=True, space_after=8)
        for pt in explain[:5]:
            _para(tf, pt, 13, INK, bullet=True, space_after=7)
    _caption(slide, s.get("caption"))
    return slide


def slide_table(prs, s):
    """A styled research table: header row in an accent color, tinted body rows.
    Spec: columns: [str,...], rows: [[str,...],...]. Optional accent: "good" |
    "warn" | "accent" (default accent), col_widths: [inches,...], caption.
    Good for 'technique | source paper | why' decision tables."""
    slide = _blank(prs)
    _bg(slide)
    _title_line(slide, s["title"])
    cols = s.get("columns", [])
    rows = s.get("rows", [])
    ncol, nrow = len(cols), len(rows) + 1
    if ncol == 0:
        return slide
    accent = {"good": GOOD, "warn": WARN}.get(s.get("accent"), ACCENT)
    tint = {"good": RGBColor(0xEC, 0xF6, 0xEF), "warn": RGBColor(0xFB, 0xEE, 0xE8)}.get(
        s.get("accent"), RGBColor(0xEE, 0xF4, 0xF3))
    left, top = Inches(0.7), Inches(1.95)
    width = Inches(11.93)
    row_h = min(Inches(0.95), (CONTENT_BOTTOM - Inches(2.05)) / max(1, nrow))
    height = int(row_h * nrow)
    gtbl = slide.shapes.add_table(nrow, ncol, int(left), int(top), int(width), int(height)).table
    gtbl.first_row = False
    gtbl.horz_banding = False
    cw = s.get("col_widths")
    if cw and len(cw) == ncol:
        for j, w in enumerate(cw):
            gtbl.columns[j].width = Inches(w)
    for j, ctext in enumerate(cols):
        cell = gtbl.cell(0, j)
        cell.fill.solid(); cell.fill.fore_color.rgb = accent
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = cell.margin_right = Inches(0.12)
        cell.margin_top = cell.margin_bottom = Inches(0.05)
        p = cell.text_frame.paragraphs[0]
        r = p.add_run(); r.text = ctext
        r.font.name = FONT; r.font.size = Pt(13); r.font.bold = True
        r.font.color.rgb = PAPER
    for i, row in enumerate(rows, start=1):
        for j in range(ncol):
            val = row[j] if j < len(row) else ""
            cell = gtbl.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = tint if (i % 2 == 1) else PAPER
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = cell.margin_right = Inches(0.12)
            cell.margin_top = cell.margin_bottom = Inches(0.05)
            p = cell.text_frame.paragraphs[0]
            r = p.add_run(); r.text = val
            r.font.name = FONT; r.font.size = Pt(12)
            r.font.bold = (j == 0)
            r.font.color.rgb = INK if j != 0 else ACCENT_DK
    _caption(slide, s.get("caption"))
    return slide


def slide_metrics(prs, s):
    """A row of big-number KPI tiles: the headline results of a progress report.
    Spec: items: [{value, label, note?}, ...] (2..4). Lead a status deck with this.
    `value` is a short string ("85%", "3.5x", "6/10"); `label` is what it measures;
    optional `note` is a small qualifier ("vs 88% best"). Optional caption."""
    slide = _blank(prs)
    _bg(slide)
    _title_line(slide, s["title"])
    items = s.get("items", [])[:4]
    n = len(items)
    if n == 0:
        return slide
    top, tile_h = Inches(2.5), Inches(3.1)
    margin, gap = Inches(0.7), Inches(0.4)
    total_w = int(EMU_W - int(margin) * 2)
    tile_w = int((total_w - int(gap) * (n - 1)) / n)
    x = int(margin)
    for it in items:
        _panel(slide, x, int(top), tile_w, int(tile_h), fill=NODE, shadow=True)
        vtf = _tb(slide, x, int(top) + Inches(0.45), tile_w, Inches(1.5),
                  anchor=MSO_ANCHOR.MIDDLE)
        _para(vtf, str(it.get("value", "")), 54, ACCENT_DK, bold=True,
              align=PP_ALIGN.CENTER, first=True, space_after=0)
        ltf = _tb(slide, x + Inches(0.25), int(top) + Inches(1.95),
                  tile_w - Inches(0.5), Inches(0.9))
        _para(ltf, str(it.get("label", "")), 17, INK, bold=True,
              align=PP_ALIGN.CENTER, first=True, space_after=2)
        if it.get("note"):
            _para(ltf, str(it["note"]), 13, MUTED, align=PP_ALIGN.CENTER,
                  space_after=0)
        x += tile_w + int(gap)
    _caption(slide, s.get("caption"))
    return slide


def slide_status_list(prs, s):
    """A list where each item carries a color-coded status badge. The workhorse
    for a 'techniques tried' slide (kept / dropped / testing) or a status roadmap.
    Spec: items: [{text, status, note?}, ...] (<=6). status: kept|done|good ->
    green, dropped|cut|bad -> rust, anything else (testing|wip|next) -> amber.
    Optional `note` prints in muted text on the right. Optional caption."""
    slide = _blank(prs)
    _bg(slide)
    _title_line(slide, s["title"])
    items = s.get("items", [])[:6]
    n = len(items)
    if n == 0:
        return slide
    top = Inches(2.0)
    left = Inches(0.7)
    width = int(EMU_W - int(left) * 2)
    gap = Inches(0.22)
    avail_h = int(CONTENT_BOTTOM - int(top))
    row_h = int((avail_h - int(gap) * (n - 1)) / n)
    pill_w = int(Inches(1.55))
    for i, it in enumerate(items):
        y = int(top) + (row_h + int(gap)) * i
        solid, tint, badge = _status_colors(it.get("status"))
        _panel(slide, int(left), y, width, row_h, fill=tint, shadow=True)
        # colored status pill on the left
        pill_h = min(int(Inches(0.5)), int(row_h * 0.6))
        py = y + (row_h - pill_h) // 2
        _rect(slide, MSO_SHAPE.ROUNDED_RECTANGLE, int(left) + Inches(0.25), py,
              pill_w, pill_h, solid)
        ptf = _tb(slide, int(left) + Inches(0.25), py, pill_w, pill_h,
                  anchor=MSO_ANCHOR.MIDDLE)
        _para(ptf, badge, 13, PAPER, bold=True, align=PP_ALIGN.CENTER,
              first=True, space_after=0)
        # main text
        tx = int(left) + Inches(0.25) + pill_w + Inches(0.3)
        note = it.get("note")
        text_w = width - (int(Inches(0.25)) + pill_w + int(Inches(0.3))) - \
            (int(Inches(3.3)) if note else int(Inches(0.3)))
        ttf = _tb(slide, tx, y, text_w, row_h, anchor=MSO_ANCHOR.MIDDLE)
        _para(ttf, str(it.get("text", "")), 18, INK, bold=True, first=True,
              space_after=0)
        if note:
            ntf = _tb(slide, int(left) + width - Inches(3.4), y, Inches(3.1),
                      row_h, anchor=MSO_ANCHOR.MIDDLE)
            _para(ntf, str(note), 15, MUTED, align=PP_ALIGN.RIGHT, first=True,
                  space_after=0)
    _caption(slide, s.get("caption"))
    return slide


def slide_timeline(prs, s):
    """A horizontal milestone timeline with status-colored dots: the 'progress so
    far, and what's next' narrative. Spec: milestones: [{label, status?}, ...]
    (<=6). status done -> green dot, now|current -> teal dot (larger), next|todo
    -> hollow/muted dot. Labels alternate above/below the line. Optional caption."""
    slide = _blank(prs)
    _bg(slide)
    _title_line(slide, s["title"])
    ms = s.get("milestones", [])[:6]
    n = len(ms)
    if n == 0:
        return slide
    line_y = int(Inches(3.9))
    margin = int(Inches(1.2))
    span = int(EMU_W) - margin * 2
    # the connecting rule
    _rect(slide, MSO_SHAPE.RECTANGLE, margin, line_y - int(Pt(1.5)),
          span, int(Pt(3)), FAINT)
    step = span / (n - 1) if n > 1 else 0
    for i, m in enumerate(ms):
        cx = margin + int(step * i) if n > 1 else margin + span // 2
        st = (m.get("status") or "").strip().lower()
        if st in ("done", "kept", "good"):
            dot, r = GOOD, int(Inches(0.16))
        elif st in ("now", "current", "active", "wip", "testing"):
            dot, r = ACCENT, int(Inches(0.21))
        else:
            dot, r = BAR, int(Inches(0.14))
        _rect(slide, MSO_SHAPE.OVAL, cx - r, line_y - r, r * 2, r * 2, dot)
        above = (i % 2 == 0)
        ly = line_y - int(Inches(1.5)) if above else line_y + int(Inches(0.45))
        anchor = MSO_ANCHOR.BOTTOM if above else MSO_ANCHOR.TOP
        ltf = _tb(slide, cx - int(Inches(1.1)), ly, int(Inches(2.2)),
                  int(Inches(1.05)), anchor=anchor)
        emph = st in ("now", "current", "active")
        _para(ltf, str(m.get("label", "")), 15,
              ACCENT_DK if emph else INK, bold=emph, align=PP_ALIGN.CENTER,
              first=True, space_after=0)
    _caption(slide, s.get("caption"))
    return slide


RENDERERS = {
    "title": slide_title,
    "table": slide_table,
    "metrics": slide_metrics,
    "status_list": slide_status_list,
    "timeline": slide_timeline,
    "summary": slide_summary,
    "section": slide_section,
    "bullets": slide_bullets,
    "two_column": slide_two_column,
    "closing": slide_closing,
    "flow": slide_flow,
    "compare": slide_compare,
    "stack": slide_stack,
    "mermaid": slide_mermaid,
}

# Full-bleed / opener / closer slides carry no footer.
FOOTER_EXEMPT = {"title", "section", "closing"}


def main():
    if len(sys.argv) != 3:
        sys.exit("usage: build_deck.py <deck.json> <out.pptx>")
    with open(sys.argv[1]) as f:
        spec = json.load(f)

    apply_palette(spec.get("theme"))

    prs = Presentation()
    prs.slide_width, prs.slide_height = EMU_W, EMU_H

    slides = spec.get("slides", [])
    total = len(slides)
    deck_title = spec.get("title") or spec.get("footer")
    if not deck_title:
        for s in slides:
            if s.get("type") == "title":
                deck_title = s.get("title")
                break

    for i, s in enumerate(slides, 1):
        r = RENDERERS.get(s.get("type"))
        if not r:
            sys.exit("unknown slide type: %s" % s.get("type"))
        slide = r(prs, s)
        if s.get("type") not in FOOTER_EXEMPT:
            _footer(slide, {"title": deck_title, "index": i, "total": total})

    prs.save(sys.argv[2])
    print("wrote %s (%d slides)" % (sys.argv[2], total))


if __name__ == "__main__":
    main()
